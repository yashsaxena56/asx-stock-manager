
import streamlit as st
import pandas as pd
from datetime import date
from io import BytesIO
from openpyxl.styles import PatternFill, Font

st.set_page_config(
    page_title="AI Stock & Expiry Manager",
    page_icon="📦",
    layout="wide"
)

st.title("📦 AI Stock & Expiry Manager")
st.write("Upload your stock Excel file to begin expiry analysis.")

uploaded_file = st.file_uploader(
    "Upload Excel or CSV file",
    type=["xlsx", "csv"]
)

if uploaded_file is not None:
    try:
        if uploaded_file.name.lower().endswith(".csv"):
            df = pd.read_csv(uploaded_file)
        else:
            df = pd.read_excel(uploaded_file)

        st.success("File uploaded successfully!")

        st.write("### Data Preview")
        st.dataframe(df.head(10), use_container_width=True)

        st.write("### File Information")
        st.write(f"Total rows: {len(df)}")
        st.write(f"Total columns: {len(df.columns)}")
        st.write("### Available Columns")
        st.write(list(df.columns))

        # STEP 1: Check required columns
        required = ["Expiry Date", "Expired Qty in Cases"]
        missing = [col for col in required if col not in df.columns]

        if missing:
            st.error(
                "Missing required columns: " + ", ".join(missing)
            )
            st.info(
                "Please check the exact column names in your file."
            )

        else:
            # STEP 2: Analyse dates and quantities
            result = df.copy()

            expiry = pd.to_datetime(
                result["Expiry Date"],
                errors="coerce",
                dayfirst=True
            )

            today = pd.Timestamp(date.today())
            days_left = (expiry.dt.normalize() - today).dt.days

            qty = pd.to_numeric(
                result["Expired Qty in Cases"],
                errors="coerce"
            )

            result["Days to Expiry"] = days_left
            result["Non-zero Expired Qty"] = (
                qty.notna() & qty.ne(0)
            )

            result["Analysis Status"] = "OK"
            result.loc[
                expiry.isna(), "Analysis Status"
            ] = "Missing/Invalid Expiry Date"
            result.loc[
                days_left < 0, "Analysis Status"
            ] = "Expired"
            result.loc[
                days_left.between(0, 30), "Analysis Status"
            ] = "Near Expiry"

            # STEP 3: Summary
            expired_count = int((days_left < 0).sum())
            near_count = int(days_left.between(0, 30).sum())
            qty_count = int(result["Non-zero Expired Qty"].sum())
            invalid_count = int(expiry.isna().sum())

            c1, c2, c3, c4 = st.columns(4)
            c1.metric("Expired Products", expired_count)
            c2.metric("Near Expiry (30 days)", near_count)
            c3.metric("Non-zero Expired Qty", qty_count)
            c4.metric("Invalid/Missing Dates", invalid_count)

            st.write("## Expiry Analysis Results")
            st.dataframe(result, use_container_width=True)

            # STEP 4: Select rows needing attention
            flagged_mask = (
                (days_left < 0)
                | days_left.between(0, 30)
                | result["Non-zero Expired Qty"]
                | expiry.isna()
            )

            flagged = result.loc[flagged_mask].copy()

            st.write("### Records Requiring Attention")
            st.dataframe(flagged, use_container_width=True)
            st.write(f"Records flagged: {len(flagged)}")

            # STEP 5: Create highlighted Excel report
            output = BytesIO()

            with pd.ExcelWriter(
                output, engine="openpyxl"
            ) as writer:
                result.to_excel(
                    writer,
                    sheet_name="Stock Analysis",
                    index=False
                )
                flagged.to_excel(
                    writer,
                    sheet_name="Flagged Records",
                    index=False
                )

                workbook = writer.book
                red = PatternFill(
                    fill_type="solid", fgColor="FFC7CE"
                )
                yellow = PatternFill(
                    fill_type="solid", fgColor="FFEB9C"
                )
                orange = PatternFill(
                    fill_type="solid", fgColor="F4B183"
                )

                for sheet in workbook.worksheets:
                    sheet.freeze_panes = "A2"
                    sheet.auto_filter.ref = sheet.dimensions

                    for cell in sheet[1]:
                        cell.font = Font(bold=True)

                    headers = {
                        cell.value: cell.column
                        for cell in sheet[1]
                    }

                    status_col = headers.get("Analysis Status")
                    qty_col = headers.get("Non-zero Expired Qty")

                    for row_num in range(2, sheet.max_row + 1):
                        status = sheet.cell(
                            row=row_num, column=status_col
                        ).value
                        qty_flag = sheet.cell(
                            row=row_num, column=qty_col
                        ).value

                        fill = None

                        if status == "Expired":
                            fill = red
                        elif status == "Near Expiry":
                            fill = yellow
                        elif (
                            status == "Missing/Invalid Expiry Date"
                            or qty_flag
                        ):
                            fill = orange

                        if fill:
                            for cell in sheet[row_num]:
                                cell.fill = fill

            st.download_button(
                label="Download Excel Analysis Report",
                data=output.getvalue(),
                file_name="stock_expiry_analysis.xlsx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                )
            )

            st.success(
                "Analysis complete! Download the report to review "
                "highlighted rows and flagged records."
            )

    except Exception as e:
        st.error(f"Could not analyse the file: {e}")

# Footer
st.markdown(
    """
    <div style="
        position: fixed;
        bottom: 10px;
        right: 18px;
        color: #888888;
        font-size: 10px;
        opacity: 0.75;
        z-index: 999;
    ">
        Crafted with care by Yash Saxena
    </div>
    """,
    unsafe_allow_html=True
)
